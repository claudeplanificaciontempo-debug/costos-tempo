# -*- coding: utf-8 -*-
import json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
filas=json.load(open('filas.json'))
VINO='9C3B4E'
def fmt(ws,n,nr,an):
    for c in range(1,n+1):
        cel=ws.cell(1,c); cel.font=Font(bold=True,color='FFFFFF',size=10)
        cel.fill=PatternFill('solid',fgColor=VINO); cel.alignment=Alignment(vertical='center',wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width=an[c-1]
    ws.freeze_panes='A2'; ws.row_dimensions[1].height=26
    ws.auto_filter.ref=f'A1:{get_column_letter(n)}{nr+1}'
BASE=[('IN-HILOPARACOSE',0.01),('IN-ETIQUETAS',1),('IN-INS001',1),('IN-CARTONERIA',1)]
INS={'BOTONES':'IN-BOTONES','BOTONES A TONO':'IN-BOTONES','BOTONES A TONO DE PRENDA':'IN-BOTONES',
 'FRAMILON':'IN-QUICKSTRETCH','BROCHES':'IN-BROCHE','ELASTICO':'IN-ELASTICO','ELASTICO EN PRETINA':'IN-ELASTICO',
 'ELASTICO Y CORDON':'IN-ELASTICO+IN-CORDONREDOND','ELASTICO, CORDON':'IN-ELASTICO+IN-CORDONREDOND',
 'ELASTICO CORDON':'IN-ELASTICO+IN-CORDONREDOND','CORDON Y ELASTICO':'IN-ELASTICO+IN-CORDONREDOND',
 'CORDON A TONO':'IN-CORDONREDOND','CORDON REDONDO':'IN-CORDONREDOND','CIERRE':'IN-CIERRENYLON',
 'CIERRE Y BOTON':'IN-CIERRENYLON+IN-BOTONES','CORDON A TONO Y CIERRE METALICO':'IN-CORDONREDOND+IN-CIERREMETALI',
 'CORDON A TONO, OJALILLOS':'IN-CORDONREDOND+IN-OJALILLO','CORDON, CIERRE, OJALILLOS':'IN-CORDONREDOND+IN-CIERRENYLON+IN-OJALILLO',
 'ELASTICO, OJALILLOS Y CORDON':'IN-ELASTICO+IN-OJALILLO+IN-CORDONREDOND',
 'ELASTICO POSTERIOR Y REGULADORES':'IN-ELASTICO+NUEVO-MUBG2490','FRAMILON, BROCHES':'IN-QUICKSTRETCH+IN-BROCHE',
 'BOTONES, FRAMILON, ELASTICO':'IN-BOTONES+IN-QUICKSTRETCH+IN-ELASTICO'}
CANT={'IN-BOTONES':3,'IN-QUICKSTRETCH':1.2,'IN-BROCHE':3,'IN-ELASTICO':1.0,'IN-CORDONREDOND':1.3,
 'IN-CIERRENYLON':1,'IN-CIERREMETALI':1,'IN-OJALILLO':2,'NUEVO-MUBG2490':2}
rec=[]; tiem=[]
for f in filas:
    ref=f['Referencia']
    cods=re.findall(r'(MP-[A-Z0-9]+|NUEVO-[A-Z0-9]+)', f['tela_tempo'])
    partes=re.findall(r'([a-zñáéíóú ]+?) ([\d.]+) (kg|m|u)', f['consumo'])
    for i,(desc,q,um) in enumerate(partes):
        if i < len(cods): cod, de = cods[i], (f['tela_pdf'] if i==0 else 'segunda tela del PDF')
        elif um=='u' and 'cuello' in desc: cod, de = 'MP-CUELLOSTEJID','cuello tejido (como en tus fichas)'
        elif um=='u': cod, de = 'MP-PUOSTEJIDOS','puños tejidos (como en tus fichas)'
        else: cod, de = 'MP-RIBB241', desc.strip()+' (como en tus fichas)'
        rec.append([ref,cod,float(q),0,de])
    for cod,q in BASE: rec.append([ref,cod,q,0,'insumo base TEMPO'])
    ins=f['insumos_pdf'].upper().strip()
    if ins and ins!='-':
        for cod in INS.get(ins,'').split('+'):
            if cod: rec.append([ref,cod,CANT.get(cod,1),0,'insumo que pide el PDF: '+ins])
    tiem.append([ref,'confeccion',f['minutos'],'mediana de tus fichas de '+f['Categoria']])
    if 'BORDADO' in f['decoracion']:
        tiem.append([ref,'bordado',int(re.search(r'(\d+)',f['decoracion']).group(1)),'puntadas · mediana de '+f['Categoria']])
wb=Workbook(); ws=wb.active; ws.title='RECETA'
ws.append(['Referencia','Codigo del material','Consumo','Desperdicio %','(de donde sale)'])
for r in rec: ws.append(r)
fmt(ws,5,len(rec),[14,26,12,14,44])
wb.save('CARGA_4_RECETA.xlsx')
wb=Workbook(); ws=wb.active; ws.title='TIEMPOS'
ws.append(['Referencia','Centro','Tiempo','(de donde sale)'])
for r in tiem: ws.append(r)
fmt(ws,4,len(tiem),[14,16,12,44])
wb.save('CARGA_5_TIEMPOS.xlsx')
print('receta:',len(rec),'| tiempos:',len(tiem))
